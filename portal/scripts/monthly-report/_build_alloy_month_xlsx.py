"""
Builds the monthly Alloy deck file (same layout as reports/alloy_live_vs_cached_2026-08.xlsx) from the JSON written by
_live_vs_cached.ts, optionally adding GHL bookings from _alloy_month_bookings.ts. Writes <out>.xlsx (formulas, recalc
afterwards in Excel), <out>.csv (UTF-8 BOM) and <out>.tsv.

python scripts/monthly-report/_build_alloy_month_xlsx.py <live_vs_cached.json> "Sep 2026" "September 1-30, 2026" <out base path> <pulled date YYYY-MM-DD> [bookings.json] [ghl_leads.json]
"""
import csv, json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, short_label, long_label, out_base, pulled = sys.argv[1:6]
bookings_src = sys.argv[6] if len(sys.argv) > 6 and sys.argv[6] else None
ghl_src = sys.argv[7] if len(sys.argv) > 7 and sys.argv[7] else None
# GMN Marketing Report payload inputs (optional): per location x offer data (_alloy_month_offers.ts) and the top
# creatives (_alloy_month_top_creatives.ts). With the offers file, <out>.payload.json is written for the report page.
offers_src = sys.argv[8] if len(sys.argv) > 8 and sys.argv[8] else None
creatives_src = sys.argv[9] if len(sys.argv) > 9 and sys.argv[9] else None
data = json.load(open(src, encoding="utf-8"))
bk = {r["name"]: r for r in json.load(open(bookings_src, encoding="utf-8"))["rows"]} if bookings_src else None
gl_doc = json.load(open(ghl_src, encoding="utf-8")) if ghl_src else None
gl = {r["name"]: r for r in gl_doc["rows"]} if gl_doc else None


TEST_RE = re.compile(r"\b(test|testing|tester|fake|asdf|qwerty|dummy)\b", re.I)


def ghl_count(g):
    """GHL leads for the report: new + returning leads, tests and duplicates removed (user, 2026-10-15)."""
    if g.get("status") != "ok":
        return None
    if "final_new" in g:  # pooled pull: final already includes returning leads
        # duplicates are removed only when the lead shows the same FULL name (user, 2026-10-16); phone/email-only matches are kept
        return g["final"] + sum(1 for d in g["duplicates"] if "same full name" not in d["on"])
    return g["final"] + len([x for x in g["returning"] if not TEST_RE.search(x["name"])])


# Calendar entries that are not appointments with anyone (user, 2026-10-16): staff block-offs and personal time, plus
# "Call with Chad" (Chad is GMN's coach). Only ever applied to events whose contact is NOT a tagged lead.
BLOCK_RE = re.compile(r"\b(unavailable|block(ed)?(\s*(off|out))?|blocked?\s*time|off|pto|vacation|holiday|out of office|ooo|flights?|airport|"
                      r"travel|trip|games?|tennis|school|dr\.?\s*appt|doctor|dentist|pick\s*up|volunteer|conference|florida|birthday)\b", re.I)
CHAD_RE = re.compile(r"\bcall\s*(with|w/|/w)\s*chad\b", re.I)


def calendar_class(a):
    """'chad' / 'block' for entries left out of the Calendar count, else ''."""
    if a.get("lead_tagged"):
        return ""
    title = a.get("title") or ""
    if CHAD_RE.search(title):
        return "chad"
    if BLOCK_RE.search(title) or (a.get("minutes") or 0) >= 120:
        return "block"
    return ""


def calendar_count(b):
    if b.get("status") != "ok" or "audit" not in b:
        return b.get("events_total") if b.get("status") == "ok" else None
    return sum(1 for a in b["audit"] if not calendar_class(a))


def event_label(a):
    """Non-lead event label: the meeting title, unless the title is just the person's name (then first name + last initial)."""
    title, who = (a.get("title") or "").strip(), a.get("who") or ""
    first = who.split(" ")[0] if who else ""
    words = title.split()
    if not title or (first and words and words[0].lower() == first.lower() and len(words) <= 3):
        return who
    if first and len(first) > 1:
        # "Starting Point Session /w Robyn Dubinsky" -> "Starting Point Session /w Robyn D."
        title = re.sub(rf"\b{re.escape(first)}\s+([A-Za-z])[\w'-]*", lambda m: f"{first} {m.group(1).upper()}.", title, flags=re.I)
    # any other "First Last" right after /w, w/, with, -, : (e.g. a coach-linked "SPS - Erin Zibrowski") -> "Erin Z."
    title = re.sub(r"((?:/w|w/|with|-|:|&)\s*)(?!Alloy\b)([A-Z][a-zA-Z]+)\s+(?:[A-Z]\.\s+)?([A-Za-z])[a-zA-Z'-]{2,}\b",
                   lambda m: f"{m.group(1)}{m.group(2)} {m.group(3).upper()}.", title)
    # a title that is only "First Last" -> "First L."
    title = re.sub(r"^([A-Z][a-z]+)\s+([A-Z])[a-zA-Z'-]+$", r"\1 \2.", title.strip())
    # phone numbers never go in the report
    title = re.sub(r"\+?\d?[\s.-]?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}", "", title)
    return re.sub(r"\s{2,}", " ", title).strip(" -")


def booking_note(b):
    """How the bookings number was reached for one location, from the per-event audit."""
    if b.get("status") != "ok":
        return b.get("status") or "dashboard not in bookings pull"
    if "audit" not in b:
        return ""
    cal = f"calendar '{b['ghl_calendar']}'" + (" (used instead of the Nurture-sheet calendar, which had no appointments this month)"
                                              if b.get("calendar_override") else "")

    def grouped(items):
        counts = {}
        for a in items:
            counts[event_label(a)] = counts.get(event_label(a), 0) + 1
        return ", ".join(f"{k} x{v}" if v > 1 else k for k, v in sorted(counts.items(), key=lambda kv: -kv[1]))

    blocks = [a for a in b["audit"] if calendar_class(a) == "block"]
    chad = [a for a in b["audit"] if calendar_class(a) == "chad"]
    text = f"CALENDAR COUNT: {b['events_total']} entries in the month on {cal}"
    if blocks:
        text += f"; minus {len(blocks)} block-off / personal time ({grouped(blocks)})"
    if chad:
        text += f"; minus {len(chad)} 'Call with Chad' (GMN coach)"
    text += f" = {calendar_count(b)}. BOOKINGS: from those {calendar_count(b)}"
    if b["cancelled"]:
        text += f", minus {b['cancelled']} cancelled"
    if b.get("invalid"):
        text += f", minus {b['invalid']} marked invalid"
    others = [a for a in b["audit"] if (a["reason"].startswith("contact not tagged") or a["reason"] == "no contact on the event")
              and not calendar_class(a)]
    if others:
        text += f", minus {len(others)} for people not tagged as ad leads (internal meetings, staff, untagged contacts): {grouped(others)}"
    reps = b.get("repeats") or []
    if reps:
        text += (f", minus {sum(p['appointments'] - 1 for p in reps)} repeat booking(s) by the same person: "
                 + ", ".join(f"{p['name']} x{p['appointments']}" for p in reps))
    text += (f" = {b['booked_ex_cancelled']} people booked ({b['showed']} showed, {b['confirmed']} confirmed/not yet marked, "
             f"{b['noshow']} no-show).")
    return text


def lead_note_pooled(g):
    """Note for the pooled pull: new + returning in one list, then tests and duplicates removed."""
    via = g.get("tag_variants") or {}
    text = (f"{g['tagged']} new contacts added in the month with tag '{g['leads_tag']}'"
            + (" or an offer version of it (" + ", ".join(f"'{k}' x{v}" for k, v in via.items()) + ")" if via else "")
            + f" (of {g['contacts_added']} contacts added) + {g['returning_found']} returning lead(s) (added earlier, filled in a form again this month)"
            + f" = {g['tagged'] + g['returning_found']}")
    if g["tests"]:
        text += f"; minus {len(g['tests'])} test: " + ", ".join(f"{t['name']} [{t['id']}]" for t in g["tests"])
    name_dups = [d for d in g["duplicates"] if "same full name" in d["on"]]
    kept_dups = [d for d in g["duplicates"] if "same full name" not in d["on"]]
    if name_dups:
        text += f"; minus {len(name_dups)} duplicate(s) with the same full name: " + "; ".join(
            f"{d['name']} [{d['id']}, {d['kind']}] = {d['on']} as {d['dup_of_name']} [{d['dup_of_id']}, {d['dup_of_kind']}, kept]" for d in name_dups)
    else:
        text += "; no duplicates (same full name)"
    if kept_dups:
        text += "; kept although phone/email matched another lead (names differ or are incomplete): " + ", ".join(
            f"{d['name']} [{d['id']}] vs {d['dup_of_name']} [{d['dup_of_id']}]" for d in kept_dups)
    new_n = g["final_new"] + sum(1 for d in kept_dups if d["kind"] == "new")
    ret_n = g["final_returning"] + sum(1 for d in kept_dups if d["kind"] == "returning")
    text += f" = {ghl_count(g)} ({new_n} new + {ret_n} returning)."
    if g.get("other_tags"):
        text += " Not counted: contacts tagged " + ", ".join(f"'{k}' x{v}" for k, v in g["other_tags"].items()) + " (a different lead tag)."
    return text


def lead_note(g):
    """How the GHL lead count was reached, in words, for one location."""
    if g.get("status") != "ok":
        return g.get("status", "")
    if "final_new" in g:
        return lead_note_pooled(g)
    via = g.get("tag_variants") or {}
    parts = [f"{g['tagged']} contacts added in the month with tag '{g['leads_tag']}'"
             + (" or an offer version of it (" + ", ".join(f"'{k}' x{v}" for k, v in via.items()) + ")" if via else "")
             + f" (of {g['contacts_added']} contacts added)"]
    if g["tests"]:
        parts.append(f"minus {len(g['tests'])} test: " + ", ".join(f"{t['name']} [{t['id']}]" for t in g["tests"]))
    if g["duplicates"]:
        parts.append(f"minus {len(g['duplicates'])} duplicate: " + ", ".join(
            f"{d['name']} [{d['id']}] = {d['on']} as {d['dup_of_name']} [{d['dup_of_id']}]" for d in g["duplicates"]))
    ret = [x for x in g["returning"] if not TEST_RE.search(x["name"])]
    ret_tests = [x for x in g["returning"] if TEST_RE.search(x["name"])]
    text = "; ".join(parts) + f" = {g['final']} new"
    if ret:
        text += f"; plus {len(ret)} returning lead(s) (added before the month, filled in a form again during it)"
    if ret_tests:
        text += f"; returning test contact(s) not counted: " + ", ".join(f"{x['name']} [{x['id']}]" for x in ret_tests)
    text += f" = {ghl_count(g)} total."
    extra = []
    if g.get("other_tags"):
        extra.append("contacts tagged " + ", ".join(f"'{k}' x{v}" for k, v in g["other_tags"].items()) + " (a different lead tag)")
    if extra:
        text += " Not counted: " + "; ".join(extra) + "."
    return text
ACCOUNT_LABELS = {
    "1007481743408819": "1007481743408819 (Non AF GMN 3 / Alloy BM1)",
    "3023829647843614": "3023829647843614 (Alloy Personal Training / BM2)",
}

rows = []
# An inactive duplicate of an active location (same campaign filter, e.g. the second "Loveland" row) would double-count it.
active_filters = {c["filter"].lower() for c in data["clients"] if c.get("active") and not c.get("forcedInactive") and c.get("filter")}
for c in data["clients"]:
    if c.get("active") is False and not c.get("forcedInactive") and (c.get("filter") or "").lower() in active_filters:
        continue
    # Alloy Ops is the rollup; test rows have an empty filter that matches every campaign.
    if c["name"].lower().startswith("alloy ops") or not c.get("filter") or " test " in f" {c['name'].lower()} ":
        continue
    L, C = c["live"], c["cached"]
    note_bits = []
    # The bookings pull lists only currently-active dashboards, so a location there was reactivated after the Meta pull.
    if (c.get("forcedInactive") or c.get("active") is False) and not (bk and c["name"] in bk):
        note_bits.append("dashboard inactive (campaign still spending)")
    if c.get("sheetLeads") is None or (c.get("sheetNote") or "").startswith("failed tabs"):
        note_bits.append("sheet tab not found")
    elif c.get("sheetNote"):
        note_bits.append(c["sheetNote"])
    if c.get("liveError"):
        note_bits.append("live error: " + c["liveError"])
    b = (bk or {}).get(c["name"]) or {}
    if bk is not None:
        if b.get("status") != "ok":
            note_bits.append("bookings: " + (b.get("status") or "dashboard not in bookings pull"))
    leads = None if "sheet tab not found" in note_bits else c.get("sheetLeads")
    ok = b.get("status") == "ok"
    g = (gl or {}).get(c["name"]) or {}
    gm = g.get("meta") or {}
    gg = g.get("ghl") or {"status": "dashboard not in GHL pull"}
    rows.append({
        "name": c["name"], "spend_l": round(L["spend"], 2), "spend_c": round(C["spend"], 2),
        "reach_u": L["reach"], "reach_c": C["reachDaily"], "reach_d": L["reachDaily"],
        "imp_l": L["impressions"], "imp_c": C["impressions"], "clk_l": L["linkClicks"], "clk_c": C["linkClicks"],
        "leads": leads, "note": "; ".join(note_bits), "accounts": c.get("accounts", []),
        "bookings": b["booked_ex_cancelled"] if ok else None, "showed": b["showed"] if ok else None,
        "noshow": b["noshow"] if ok else None, "cancelled": b["cancelled"] if ok else None,
        "booking_note": booking_note(b) if bk is not None else "", "bk": b,
        "calendar_count": calendar_count(b) if ok else None,
        "meta_status": gm.get("status"), "meta_spend7": gm.get("spend_7d"), "meta_last_spend": gm.get("last_spend"),
        "ghl_leads": ghl_count(gg), "ghl_note": lead_note(gg) if gl is not None else "",
        "ghl": gg,
    })
# Locations with no ad spend in the month are left out of the report (user, 2026-10-15).
dropped_no_spend = [r["name"] for r in rows if r["spend_l"] == 0 and r["spend_c"] == 0]
rows = [r for r in rows if r["name"] not in dropped_no_spend]
rows.sort(key=lambda r: r["name"])
ops = next((c for c in data["clients"] if c["name"].lower().startswith("alloy ops")), None)

MONEY, INT, PCT = "\\$#,##0.00", "#,##0", "0.00%"
HEAD = ["Location", "Spend (live)", "Spend (cached)", "Reach (live, Meta unique)", "Reach (cached dashboard)",
        "Reach (live, daily-sum method)", "Impr. (live)", "Impr. (cached)", "Link clicks (live)", "Link clicks (cached)",
        "CTR (live)", "CTR (cached)", "Leads (sheet)", "CPL (live)", "CPL (cached)"]
BOOK_HEAD = ["Bookings (GHL)", "Showed", "No-shows", "Cancelled (not counted)", "Cost per booking",
             "Calendar count (all appointments shown in GHL)", "Cost per calendar booking", "Booking notes"] if bk is not None else []
meta_asof = gl_doc["meta_window"].split("..")[1] if gl_doc else ""
GHL_HEAD = [f"Active on Meta (as of {meta_asof})", "Meta spend, last 7 days", "Last Meta spend date", "GHL leads", "GHL lead notes"] if gl is not None else []
GHL_START = len(HEAD) + len(BOOK_HEAD) + 1  # 1-based column of the first GHL/Meta-status column
HEAD = HEAD + BOOK_HEAD + GHL_HEAD + ["Note"]
NOTE_COL = len(HEAD)
LASTCOL = get_column_letter(NOTE_COL)
NAVY = PatternFill("solid", fgColor="FF1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="FFD9E1F2")
WHITE_B = Font(b=True, color="FFFFFFFF")

wb = Workbook()
summary = wb.active
summary.title = "Summary"
sheet_name = f"Alloy {short_label}"
ws = wb.create_sheet(sheet_name)

for i, h in enumerate(HEAD, 1):
    cell = ws.cell(1, i, h)
    cell.font, cell.fill = WHITE_B, NAVY
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
ws.row_dimensions[1].height = 46
last = len(rows) + 1
t = last + 1
for r_i, r in enumerate(rows, 2):
    vals = [r["name"], r["spend_l"], r["spend_c"], r["reach_u"], r["reach_c"], r["reach_d"], r["imp_l"], r["imp_c"],
            r["clk_l"], r["clk_c"], f'=IF(G{r_i}>0,I{r_i}/G{r_i},"")', f'=IF(H{r_i}>0,J{r_i}/H{r_i},"")',
            r["leads"], f'=IF(N(M{r_i})>0,B{r_i}/M{r_i},"")', f'=IF(N(M{r_i})>0,C{r_i}/M{r_i},"")']
    fmts = [None, MONEY, MONEY, INT, INT, INT, INT, INT, INT, INT, PCT, PCT, INT, MONEY, MONEY]
    if bk is not None:
        vals += [r["bookings"], r["showed"], r["noshow"], r["cancelled"], f'=IF(N(P{r_i})>0,B{r_i}/P{r_i},"")',
                 r["calendar_count"], f'=IF(N(U{r_i})>0,B{r_i}/U{r_i},"")', r["booking_note"] or None]
        fmts += [INT, INT, INT, INT, MONEY, INT, MONEY, None]
    if gl is not None:
        vals += [r["meta_status"], r["meta_spend7"], r["meta_last_spend"], r["ghl_leads"], r["ghl_note"] or None]
        fmts += [None, MONEY, None, INT, None]
    vals.append(r["note"] or None)
    fmts.append(None)
    for c_i, (v, f) in enumerate(zip(vals, fmts), 1):
        cell = ws.cell(r_i, c_i, v)
        if f:
            cell.number_format = f
tot = ["TOTAL (all Alloy locations)"] + [f"=SUM({col}2:{col}{last})" for col in "BCDEFGHIJ"] + \
      [f'=IF(G{t}>0,I{t}/G{t},"")', f'=IF(H{t}>0,J{t}/H{t},"")', f"=SUM(M2:M{last})",
       f'=IF(M{t}>0,B{t}/M{t},"")', f'=IF(M{t}>0,C{t}/M{t},"")']
tfmt = [None, MONEY, MONEY, INT, INT, INT, INT, INT, INT, INT, PCT, PCT, INT, MONEY, MONEY]
if bk is not None:
    # Cost per booking only over locations that have booking data, so unconnected locations' spend doesn't inflate it.
    tot += [f"=SUM({col}2:{col}{last})" for col in "PQRS"] + [f'=IF(P{t}>0,SUMIF(P2:P{last},"<>",B2:B{last})/P{t},"")',
                                                              f"=SUM(U2:U{last})", f'=IF(U{t}>0,SUMIF(U2:U{last},"<>",B2:B{last})/U{t},"")', None]
    tfmt += [INT, INT, INT, INT, MONEY, INT, MONEY, None]
if gl is not None:
    sc, s7, gc = (get_column_letter(GHL_START + i) for i in (0, 1, 3))
    tot += [f'=COUNTIF({sc}2:{sc}{last},"Active")&" of "&COUNTA(A2:A{last})&" active"', f"=SUM({s7}2:{s7}{last})", None,
            f"=SUM({gc}2:{gc}{last})", None]
    tfmt += [None, MONEY, None, INT, None]
tot.append("Reach totals add locations together; a person reached by two locations is counted twice")
tfmt.append(None)
for c_i, (v, f) in enumerate(zip(tot, tfmt), 1):
    cell = ws.cell(t, c_i, v)
    cell.font, cell.fill = Font(b=True), TOTAL_FILL
    if f:
        cell.number_format = f
widths = {"A": 46, "B": 14, "D": 15, "F": 16, "G": 14, "I": 12, "K": 10, "M": 11, LASTCOL: 60}
if bk is not None:
    widths.update({"P": 11, "Q": 9, "R": 9, "S": 11, "T": 12, "U": 14, "V": 12, "W": 90})
if gl is not None:
    for i, w in enumerate([22, 12, 12, 10, 90]):
        widths[get_column_letter(GHL_START + i)] = w
for col, w in widths.items():
    ws.column_dimensions[col].width = w
# "(live)" columns are hidden, not removed (user, 2026-10-16): formulas such as cost per booking still read them.
for i, h in enumerate(HEAD, 1):
    if "(live" in h:
        ws.column_dimensions[get_column_letter(i)].hidden = True
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{LASTCOL}{last}"

# Summary
q = f"'{sheet_name}'"
summary["A1"] = f"Alloy - {long_label}"
summary["A1"].font = Font(b=True, sz=14)
for col, h in zip("BCD", ["Meta live", "Cached dashboard", "Difference"]):
    summary[f"{col}2"] = h
    summary[f"{col}2"].font, summary[f"{col}2"].fill = WHITE_B, NAVY
lines = [
    ("Spend", "B", "C", MONEY), ("Reach (Meta unique vs dashboard daily-sum)", "D", "E", INT),
    ("Reach, same daily-sum method on both sides", "F", "E", INT), ("Impr.", "G", "H", INT),
    ("Link clicks", "I", "J", INT), ("CTR", "K", "L", PCT), ("Leads (sheet)", "M", "M", INT), ("CPL", "N", "O", MONEY),
]
for i, (label, lc, cc, f) in enumerate(lines, 3):
    summary[f"A{i}"] = label
    summary[f"A{i}"].font = Font(b=True)
    summary[f"B{i}"] = f"={q}!{lc}{t}"
    summary[f"C{i}"] = f"={q}!{cc}{t}"
    summary[f"D{i}"] = f"=C{i}-B{i}"
    for col in "BCD":
        summary[f"{col}{i}"].number_format = f
stats = [
    ("Total locations", f"=COUNTA({q}!A2:A{last})", INT),
    ("Locations with ad spend", f'=COUNTIF({q}!B2:B{last},">0")', INT),
    ("Avg. location spend (locations with spend)", '=IF(B13>0,B3/B13,"")', MONEY),
    ("Avg. location spend (all locations)", '=IF(B12>0,B3/B12,"")', MONEY),
]
if bk is not None:
    stats += [
        ("Bookings (GHL)", f"={q}!P{t}", INT),
        ("Calendar count (all appointments shown in GHL)", f"={q}!U{t}", INT),
        ("Cost per calendar booking (locations with booking data)", f"={q}!V{t}", MONEY),
        ("Locations with booking data", f'=COUNT({q}!P2:P{last})', INT),
        ("Cost per booking (locations with booking data)", f"={q}!T{t}", MONEY),
    ]
if gl is not None:
    sc, gc = get_column_letter(GHL_START), get_column_letter(GHL_START + 3)
    stats += [
        (f"Locations active on Meta (as of {meta_asof})", f'=COUNTIF({q}!{sc}2:{sc}{last},"Active")', INT),
        ("GHL leads (new + returning, tests and duplicates removed)", f"={q}!{gc}{t}", INT),
        ("Locations with GHL lead data", f"=COUNT({q}!{gc}2:{gc}{last})", INT),
    ]
for i, (label, formula, f) in enumerate(stats, 12):
    summary[f"A{i}"] = label
    summary[f"A{i}"].font = Font(b=True)
    summary[f"B{i}"] = formula
    summary[f"B{i}"].number_format = f
notes_head = 12 + len(stats) + 2

reach_u = sum(r["reach_u"] for r in rows)
reach_d = sum(r["reach_d"] for r in rows)
ratio = reach_d / reach_u if reach_u else 0
reach_match = all(r["reach_d"] == r["reach_c"] for r in rows)
accts = sorted({a for r in rows for a in r["accounts"]})
acct_spend = {t_["account"]: (t_["live"] or {}).get("spend", 0) for t_ in data.get("accountTotals", [])}
inactive = [r["name"] for r in rows if "dashboard inactive" in r["note"]]
no_tab = [r for r in rows if "sheet tab not found" in r["note"]]
notes = [
    f"Live: Meta Marketing API, campaign level, pulled {pulled} for ad accounts "
    + " and ".join(f"{ACCOUNT_LABELS.get(a, a)}, ${acct_spend.get(a, 0):,.2f} account spend this month" for a in accts)
    + ". Cached: Marketing Dashboard database (meta_daily_insights). Both use each location's own dashboard campaign filter.",
    "Reach is a count of unique people, so it cannot be added up like spend. 'Reach (live, Meta unique)' is Meta's unique reach "
    "per campaign for the whole month, added across the location's campaigns. The dashboard instead adds each day's reach, which "
    f"counts the same person again on every day they saw an ad, so it reads about {ratio:.1f}x higher."
    + (f" Counted the same daily-sum way, live and cached match exactly for all {len(rows)} locations, so this is a counting-method "
       "difference, not a sync error." if reach_match else ""),
    "Reach totals add locations (and campaigns) together, so someone reached by two campaigns is counted twice. "
    "Meta Ads Manager's own account-level total would be lower.",
    f"CTR = link clicks / impressions. Leads come from each location's tab in the Alloy Daily FB Data Google Sheet, {long_label}; "
    "the same lead count is used for both CPL columns. CPL = spend / sheet leads.",
]
if bk is not None:
    no_bk = [r["name"] for r in rows if r["bookings"] is None]
    notes.append(
        f"Bookings = distinct PEOPLE tagged as ad leads (the location's Leads Tag or an offer version of it) who had at least one "
        f"appointment scheduled {long_label} on the location's booking calendar in GoHighLevel (the calendar listed in the Nurture "
        "Team booking sheet) that was not cancelled. No-shows count, because the person did book. Appointments for contacts without "
        "the ad-lead tag (internal meetings such as account-manager calls, staff time blocks, vendors, untagged walk-ins) are left out, "
        "and a person who booked more than once counts once. Showed / No-shows count people by their best status in the month "
        "(a no-show who later showed counts as showed); Cancelled counts cancelled appointments. 'Booking notes' shows the steps for "
        "each location and the 'Booking audit' sheet lists every appointment with the reason it was or was not counted. "
        "Cost per booking = spend / bookings; the total uses only the spend of locations that have booking data. "
        "'Calendar count' is the number of appointments on that same calendar in the month as GoHighLevel's calendar view shows "
        "them, minus staff block-offs and personal time (Unavailable, flights, games, school, doctor appointments, all-day or 2+ hour "
        "blocks) and 'Call with Chad' (GMN's coach); it still includes cancelled (crossed-out) appointments, internal meetings and "
        "people who booked more than once. Cost per calendar booking = spend / calendar count (total uses only locations with booking "
        "data). 'Booking notes' shows both steps for every location.")
    if no_bk:
        notes.append(f"Blank Bookings ({len(no_bk)} locations): no GoHighLevel connection on the dashboard yet, so their calendars "
                     "cannot be read: " + ", ".join(no_bk) + ".")
if gl is not None:
    notes.append(
        f"Active on Meta = the location has at least one campaign whose status is ACTIVE in Meta right now AND that campaign spent "
        f"money in the last 7 synced days ({gl_doc['meta_window'].replace('..', ' to ')}). 'Campaign on, $0 spend' means a campaign is "
        "switched on but did not deliver; 'No active campaign' means nothing is running. Last Meta spend date is the last day any of "
        "the location's campaigns spent.")
    notes.append(
        f"GHL leads = contacts ADDED to the location's GoHighLevel account during {long_label} (dated in the ad account's timezone, "
        f"{gl_doc['tz']}, the same as the dashboard) that carry the location's Leads Tag or an offer version of it (e.g. 'new ad lead - gmn - 2 weeks free', "
        "'new ad lead - gmn challenge'; these are GMN ad leads for another offer and are counted). Then test contacts are removed (name or email "
        "containing test / fake / asdf / dummy), then duplicates: a lead is removed only when it shows the same FULL name (first + last) "
        "as an earlier lead in the month (the earliest is kept); a matching phone or email alone does not remove anyone and is only "
        "noted. New and returning leads are checked against each other. 'GHL lead notes' shows each removal by first name + last "
        "initial and GHL contact ID, and the 'GHL lead exclusions' sheet lists every one. Returning leads are then ADDED: contacts "
        "added before the month who filled in a form again during it (the dashboard's GHL leads card counts these the same way). "
        "Not counted, but noted: contacts with a "
        "different 'new ad lead' tag that is not an offer version of the location's Leads Tag. GoHighLevel already merges contacts "
        "that share a phone or email, which is why few duplicates remain. GHL leads differ from Leads (sheet), which come from "
        "the Alloy Daily FB Data sheet.")
if dropped_no_spend:
    notes.append("Left out because they had no ad spend this month: " + ", ".join(dropped_no_spend) + ".")
if inactive:
    notes.append("Rows noted 'dashboard inactive': " + ", ".join(inactive) + " still had ad spend this month but their dashboard is "
                 "switched off. They are included so the location rows add up to the Alloy Ops total"
                 + (f" (${ops['live']['spend']:,.2f} spend, {ops['sheetLeads']:,} sheet leads)." if ops else "."))
if no_tab:
    spent = [r["name"] for r in no_tab if r["spend_l"] > 0]
    notes.append("Blank Leads with note 'sheet tab not found': " + ", ".join(r["name"] for r in no_tab) + " have no sheet tab yet"
                 + ("; all had $0 spend this month." if not spent else f"; with spend: {', '.join(spent)}."))
summary[f"A{notes_head}"] = "Notes and sources"
summary[f"A{notes_head}"].font = Font(b=True)
for i, n in enumerate(notes, notes_head + 1):
    summary[f"A{i}"] = n
    summary[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells(f"A{i}:D{i}")
    summary.row_dimensions[i].height = 62 if len(n) < 420 else 92
summary.column_dimensions["A"].width = 50
summary.column_dimensions["B"].width = 20
summary.column_dimensions["C"].width = 18
summary.column_dimensions["D"].width = 14

if bk is not None and any("audit" in (r["bk"] or {}) for r in rows):
    au = wb.create_sheet("Booking audit")
    au_head = ["Location", "Date", "Person (first name + last initial)", "GHL contact ID", "Event title (non-lead events only)",
               "Status", "Minutes", "Tagged as ad lead?", "In calendar count?", "Counted as booking?", "Reason"]
    for i, h in enumerate(au_head, 1):
        cell = au.cell(1, i, h)
        cell.font, cell.fill = WHITE_B, NAVY
    rr = 2
    for r in rows:
        for a in (r["bk"] or {}).get("audit", []):
            cls = calendar_class(a)
            reason = {"block": "block-off / personal time (left out of calendar count and bookings)",
                      "chad": "Call with Chad, GMN coach (left out of calendar count and bookings)"}.get(cls, a["reason"])
            for c_i, v in enumerate([r["name"], a["date"], a["who"], a["contact_id"], (event_label(a) if a["title"] else None), a["status"], a["minutes"],
                                     "Yes" if a["lead_tagged"] else "No", "No" if cls else "Yes", "Yes" if a["counted"] else "No", reason], 1):
                au.cell(rr, c_i, v)
            rr += 1
    for col, w in zip("ABCDEFGHIJK", [46, 11, 24, 24, 44, 11, 8, 10, 10, 10, 60]):
        au.column_dimensions[col].width = w
    au.freeze_panes = "A2"
    au.auto_filter.ref = f"A1:K{max(rr - 1, 1)}"

if gl is not None:
    ex = wb.create_sheet("GHL lead exclusions")
    ex_head = ["Location", "Lead (first name + last initial)", "GHL contact ID", "Date added", "Counted?", "Reason", "Kept instead / detail"]
    for i, h in enumerate(ex_head, 1):
        cell = ex.cell(1, i, h)
        cell.font, cell.fill = WHITE_B, NAVY
    rr = 2
    for r in rows:
        g = r["ghl"]
        if g.get("status") != "ok":
            continue
        items = [(t["name"], t["id"], t["added"][:10], "No", "Test contact" + (f" ({t['kind']})" if t.get("kind") else ""),
                  f"email domain {t['email_domain']}" if t["email_domain"] else "") for t in g["tests"]]
        items += [(d["name"], d["id"], d["added"][:10], "No" if "same full name" in d["on"] else "Yes",
                   (f"Duplicate - {d['on']}" if "same full name" in d["on"] else f"Possible duplicate ({d['on']}) - kept, names differ")
                   + (f" ({d['kind']} lead)" if d.get("kind") else ""),
                   f"kept {d['dup_of_name']} [{d['dup_of_id']}]" + (f" ({d['dup_of_kind']} lead)" if d.get("dup_of_kind") else "")) for d in g["duplicates"]]
        items += [(x["name"], x["id"], x["first_added"], "No" if TEST_RE.search(x["name"]) else "Yes",
                   "Returning lead" + (" - test contact" if TEST_RE.search(x["name"]) else ""),
                   "added before the month, filled in a form again during it") for x in g["returning"]]
        for tag, n in (g.get("tag_variants") or {}).items():
            items.append((f"{n} contact(s)", "", "", "Yes", f"Counted via offer tag '{tag}' (not the exact Leads Tag)",
                          "the dashboard's GHL card misses these until the tag is added to the Leads Tag in admin"))
        for tag, n in (g.get("other_tags") or {}).items():
            items.append((f"{n} contact(s)", "", "", "No", f"Different lead tag '{tag}'", "not an offer version of this location's Leads Tag"))
        for it in items:
            for c_i, v in enumerate((r["name"],) + it, 1):
                ex.cell(rr, c_i, v)
            rr += 1
    for col, w in zip("ABCDEFG", [46, 26, 26, 12, 10, 52, 48]):
        ex.column_dimensions[col].width = w
    ex.freeze_panes = "A2"
    ex.auto_filter.ref = f"A1:G{max(rr - 1, 1)}"
wb.save(out_base + ".xlsx")

# csv/tsv with computed values, same columns as the xlsx
def blank(v):
    return "" if v is None else v

def fmt_row(r):
    ctr_l = f"{r['clk_l'] / r['imp_l'] * 100:.2f}%" if r["imp_l"] else ""
    ctr_c = f"{r['clk_c'] / r['imp_c'] * 100:.2f}%" if r["imp_c"] else ""
    has = isinstance(r["leads"], int) and r["leads"] > 0
    out = [r["name"], f"{r['spend_l']:.2f}", f"{r['spend_c']:.2f}", r["reach_u"], r["reach_c"], r["reach_d"], r["imp_l"],
           r["imp_c"], r["clk_l"], r["clk_c"], ctr_l, ctr_c, blank(r["leads"]),
           f"{r['spend_l'] / r['leads']:.2f}" if has else "", f"{r['spend_c'] / r['leads']:.2f}" if has else ""]
    if bk is not None:
        cpb_spend = r.get("bk_spend", r["spend_l"])
        out += [blank(r["bookings"]), blank(r["showed"]), blank(r["noshow"]), blank(r["cancelled"]),
                f"{cpb_spend / r['bookings']:.2f}" if r["bookings"] else "", blank(r.get("calendar_count")),
                f"{cpb_spend / r['calendar_count']:.2f}" if r.get("calendar_count") else "", r.get("booking_note") or ""]
    if gl is not None:
        out += [blank(r.get("meta_status")), "" if r.get("meta_spend7") is None else f"{r['meta_spend7']:.2f}",
                blank(r.get("meta_last_spend")), blank(r.get("ghl_leads")), r.get("ghl_note") or ""]
    return out + [r["note"]]

total = {k: sum(r[k] for r in rows) for k in ["spend_l", "spend_c", "reach_u", "reach_c", "reach_d", "imp_l", "imp_c", "clk_l", "clk_c"]}
total.update(name="TOTAL (all Alloy locations)", leads=sum(r["leads"] or 0 for r in rows),
             note="Reach totals add locations together; a person reached by two locations is counted twice")
if bk is not None:
    for k in ["bookings", "showed", "noshow", "cancelled"]:
        total[k] = sum(r[k] or 0 for r in rows)
    total["bk_spend"] = sum(r["spend_l"] for r in rows if r["bookings"] is not None)
    total["calendar_count"] = sum(r["calendar_count"] or 0 for r in rows)
    total["booking_note"] = ""
if gl is not None:
    total.update(meta_status=f"{sum(1 for r in rows if r['meta_status'] == 'Active')} of {len(rows)} active",
                 meta_spend7=sum(r["meta_spend7"] or 0 for r in rows), meta_last_spend=None,
                 ghl_leads=sum(r["ghl_leads"] or 0 for r in rows), ghl_note="")
out_rows = [fmt_row(r) for r in rows] + [fmt_row(total)]
with open(out_base + ".csv", "w", newline="", encoding="utf-8-sig") as f:
    csv.writer(f).writerows([HEAD] + out_rows)
with open(out_base + ".tsv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f, delimiter="\t", lineterminator="\n").writerows([HEAD] + out_rows)
print(f"{len(rows)} locations; spend live {total['spend_l']:.2f} cached {total['spend_c']:.2f}; leads {total['leads']}; "
      f"reach unique {total['reach_u']} daily-sum {total['reach_d']} ({ratio:.2f}x); Ops "
      + (f"{ops['live']['spend']:.2f} / {ops['sheetLeads']}" if ops else "n/a")
      + (f"; bookings {total['bookings']} at {sum(1 for r in rows if r['bookings'] is not None)} locations, "
         f"cost/booking {total['bk_spend'] / total['bookings']:.2f}" if bk is not None and total['bookings'] else ""))


# ---------------------------------------------------------------------------------------------------------------------
# GMN Marketing Report payload (<out>.payload.json) — the same numbers as the deck, published to the report page with
# portal/_publish_monthly_report.ts. Deck mapping: cached Meta columns, Reach = dashboard daily-sum, Bookings = calendar
# count, CPB = spend / calendar count (total over locations with booking data).
if offers_src:
    od = json.load(open(offers_src, encoding="utf-8"))
    month = od["month"]

    def metric_row(label, spend, reach, imp, clk, leads, bookings, cpb_spend):
        return {
            "label": label, "spend": round(spend, 2), "reach": int(reach), "impressions": int(imp), "linkClicks": int(clk),
            "ctr": round(clk / imp * 100, 2) if imp else None, "leads": leads,
            "cpl": round(spend / leads, 2) if leads else None,
            "bookings": None if bookings is None else int(round(bookings)),
            "cpb": round(cpb_spend / bookings, 2) if bookings else None,
        }

    loc_rows = [metric_row(r["name"], r["spend_c"], r["reach_c"], r["imp_c"], r["clk_c"], r["leads"],
                           r.get("calendar_count"), r["spend_c"]) for r in rows]
    have_bk = [r for r in rows if r.get("calendar_count") is not None]
    loc_total = metric_row("Total", sum(r["spend_c"] for r in rows), sum(r["reach_c"] for r in rows),
                           sum(r["imp_c"] for r in rows), sum(r["clk_c"] for r in rows), sum(r["leads"] or 0 for r in rows),
                           sum(r["calendar_count"] for r in have_bk), sum(r["spend_c"] for r in have_bk))

    def split_int(total, weights):
        """Split an integer total by weights into integers that add back up (largest remainder)."""
        s = sum(weights)
        if not total or s <= 0:
            return [0] * len(weights)
        raw = [total * w / s for w in weights]
        out = [int(x) for x in raw]
        for i in sorted(range(len(raw)), key=lambda i: raw[i] - out[i], reverse=True)[: total - sum(out)]:
            out[i] += 1
        return out

    offers, mismatch = {}, []
    for r in rows:
        per = od["locations"].get(r["name"], {})
        if abs(sum(v["spend"] for v in per.values()) - r["spend_c"]) > 0.01:
            mismatch.append((r["name"], round(sum(v["spend"] for v in per.values()), 2), r["spend_c"]))
        names = list(per)
        cal = r.get("calendar_count")
        lead_w = [per[n]["sheet_leads"] for n in names]
        shares = split_int(cal or 0, lead_w if sum(lead_w) > 0 else [per[n]["spend"] for n in names])
        for n, b in zip(names, shares):
            o = offers.setdefault(n, {"spend": 0.0, "reach": 0, "imp": 0, "clk": 0, "leads": 0, "bookings": 0, "cpb_spend": 0.0, "has_bk": False})
            v = per[n]
            o["spend"] += v["spend"]; o["reach"] += v["reach"]; o["imp"] += v["impressions"]; o["clk"] += v["link_clicks"]
            o["leads"] += v["sheet_leads"]
            if cal is not None:
                o["bookings"] += b; o["cpb_spend"] += v["spend"]; o["has_bk"] = True
    offer_rows = sorted((metric_row(n, o["spend"], o["reach"], o["imp"], o["clk"], o["leads"],
                                    o["bookings"] if o["has_bk"] else None, o["cpb_spend"]) for n, o in offers.items()),
                        key=lambda x: -x["spend"])
    offers_total = metric_row("Total", sum(o["spend"] for o in offers.values()), sum(o["reach"] for o in offers.values()),
                              sum(o["imp"] for o in offers.values()), sum(o["clk"] for o in offers.values()),
                              sum(o["leads"] for o in offers.values()), sum(o["bookings"] for o in offers.values()),
                              sum(o["cpb_spend"] for o in offers.values()))

    creatives = json.load(open(creatives_src, encoding="utf-8")) if creatives_src else {"topCreatives": [], "note": None}
    payload = {
        "version": 1,
        "title": "Alloy Personal Training Local Marketing Monthly Report",
        "month": month,
        "monthLabel": long_label.split(" ")[0] + " " + month[:4],
        "generatedAt": pulled,
        "overview": {
            "totalLocations": len(rows),
            "spend": loc_total["spend"], "leads": loc_total["leads"], "cpl": loc_total["cpl"],
            "bookings": loc_total["bookings"], "cpb": loc_total["cpb"],
            "avgLocationSpend": round(loc_total["spend"] / len(rows), 2) if rows else None,
        },
        "locations": loc_rows, "locationsTotal": loc_total,
        "offers": offer_rows, "offersTotal": offers_total,
        "topCreatives": creatives.get("topCreatives", []),
        "creativesNote": creatives.get("note"),
        "notes": [
            "Spend, reach, impressions and link clicks: Meta (dashboard data). Reach adds each day's reach, so a person "
            "reached on several days is counted again.",
            "Leads: each location's Alloy Daily FB Data sheet. CPL = spend / leads.",
            "Bookings: appointments on each location's GoHighLevel booking calendar in the month, excluding staff block-offs "
            "and personal time; CPB = spend / bookings. In the offer breakdown, a location's bookings are split across its "
            "offers by that location's lead share.",
        ],
    }
    open(out_base + ".payload.json", "w", encoding="utf-8").write(json.dumps(payload, indent=1))
    print(f"payload: {len(loc_rows)} locations, {len(offer_rows)} offers, {len(payload['topCreatives'])} creatives; "
          f"locations total ${loc_total['spend']:,.2f} / {loc_total['leads']} leads / {loc_total['bookings']} bookings / CPB {loc_total['cpb']}; "
          f"offers total ${offers_total['spend']:,.2f} / {offers_total['leads']} leads / {offers_total['bookings']} bookings / CPB {offers_total['cpb']}")
    if mismatch:
        print("WARNING offer spend != location spend:", mismatch)
